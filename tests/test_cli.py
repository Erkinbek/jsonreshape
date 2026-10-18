"""The command line (AGENT_PROMPT §13, §11, §1.4; docs/ARCHITECTURE.md §10), in-process.

Every command and option family runs through ``cli.main(argv)`` with stdout/stderr captured,
on small generated files; outputs are verified by reading them back (codec, csv, openpyxl,
sqlite3). Covered: exit codes 0/1/2/3/4 (bad expression with a caret line, missing file,
JSON error with line/column, ``--allow-partial``, OUTPUT == INPUT, cancellation through an
injected CancelToken and through SIGINT), Uzbek and English messages (argparse's own phrases
included), ``inspect --json``, highlighting on/off, the progress line, ``--export`` round
trip against ``search.reference_search``, ``--backend python`` vs ``--backend rust``
(subprocesses: the backend of a process is fixed when its core is imported) and AGENT_PROMPT
§14 test 24 (``0.1``, ``1e400``, ``12345678901234567890``, ``-0.0``, ``1.10`` survive
filter/sort/convert to JSON and JSONL byte-identical and unquoted).
"""

from __future__ import annotations

import argparse
import ast
import csv
import errno
import io
import json
import logging
import os
import re
import signal
import sqlite3
import subprocess
import sys
import threading
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

import pytest

import jsongen
import jsonreshape
from jsonreshape import cli, codec, i18n
from jsonreshape.tasks import CancelToken

ROOT = Path(__file__).resolve().parents[1]
VERSION = jsonreshape.__version__
BACKSLASH_U = bytes([92, 117])  # the two bytes of a "\u" escape

TOSHKENT_CYR = "".join(map(chr, (0x0422, 0x043E, 0x0448, 0x043A, 0x0435, 0x043D, 0x0442)))  # Toshkent in Cyrillic
OKINA = chr(0x02BB)  # the Uzbek apostrophe (oʻ)
N = codec.loads      # exact number values (JsonFloat / JsonInt keep their token)


# ----------------------------------------------------------------------------------
# fixtures and helpers
# ----------------------------------------------------------------------------------
@pytest.fixture(autouse=True)
def _cli_isolation(monkeypatch: pytest.MonkeyPatch):
    """main() changes process state on purpose (language, JSONRESHAPE_BACKEND, logging,
    signal handlers); put everything back after each test."""
    old_backend = os.environ.get("JSONRESHAPE_BACKEND")
    monkeypatch.setenv("JSONRESHAPE_LANG", "uz")
    monkeypatch.delenv("NO_COLOR", raising=False)
    old_lang = i18n.get_language()
    i18n.set_language("uz")   # whatever an earlier test left behind
    root = logging.getLogger()
    old_level = root.level
    old_sigint = signal.getsignal(signal.SIGINT)
    # the end-of-command memory check would measure this whole pytest process (it can pass
    # the 512 MB default): off here; tested in test_fix2_cli.py
    monkeypatch.setattr(cli, "_peak_memory", lambda: None)
    yield
    if old_backend is None:
        os.environ.pop("JSONRESHAPE_BACKEND", None)
    else:
        os.environ["JSONRESHAPE_BACKEND"] = old_backend
    i18n.set_language(old_lang)
    from jsonreshape import log as applog

    applog.shutdown_logging()
    root.setLevel(old_level)
    assert signal.getsignal(signal.SIGINT) == old_sigint, "main() must restore the SIGINT handler"


@dataclass
class Result:
    code: int
    out: str
    err: str


@pytest.fixture
def run(capsys: pytest.CaptureFixture[str]):
    def _run(*argv: Any, cancel: CancelToken | None = None) -> Result:
        code = cli.main([str(a) for a in argv], cancel=cancel)
        captured = capsys.readouterr()
        return Result(code, captured.out, captured.err)

    return _run


def people() -> list[dict[str, Any]]:
    """Eight records with controlled values (filter/sort/search expectations below)."""
    return [
        {"id": 1, "name": "Anvar", "score": 10, "active": True, "tags": ["x", "y"],
         "contact": {"address": {"city": "Toshkent"}}, "orders": [{"oid": 7, "sum": N("1.10")}, {"oid": 8, "sum": N("0.1")}]},
        {"id": 2, "name": "Bobur", "score": 25, "active": False, "tags": [],
         "contact": {"address": {"city": TOSHKENT_CYR}}, "orders": []},
        {"id": 3, "name": "dilmurod", "score": None, "active": True, "tags": ["z"],
         "contact": {"address": {"city": "Samarqand"}}},
        {"id": 4, "name": "Eldor", "score": N("1e400"), "active": True,
         "contact": {"address": {"city": "O" + OKINA + "zbekiston"}},
         "orders": [{"oid": 9, "sum": N("12345678901234567890")}]},
        {"id": 5, "name": "child10", "score": N("-0.0"), "active": False, "tags": ["Toshkent shahri"]},
        {"id": 6, "name": "child2", "score": 3, "active": True},
        {"id": 7, "name": "Farrux", "score": 25, "active": True, "contact": {"address": {"city": "toshkent"}}},
        {"id": 8, "name": "aziz", "active": False},
    ]


@pytest.fixture
def src(tmp_path: Path) -> Path:
    p = tmp_path / "people.json"
    jsongen.write_file(p, people(), "array", style="pretty2")
    return p


@pytest.fixture
def rnd(tmp_path: Path) -> Path:
    """Random realistic records (escapes, Uzbek/Cyrillic text, exact number tokens)."""
    p = tmp_path / "random.json"
    jsongen.write_file(p, jsongen.records(60, seed=11), "array", style="pretty2")
    return p


def canon(value: Any) -> bytes:
    return codec.dumps(value)


def read_json(path: Path) -> Any:
    return codec.loads(path.read_bytes())


def read_jsonl(path: Path) -> list[Any]:
    data = path.read_bytes()
    assert data.endswith(b"\n")
    return [codec.loads(line) for line in data.split(b"\n") if line]


def ids(records: list[Any]) -> list[int]:
    return [r["id"] for r in records]


def read_csv(path: Path, delimiter: str = ",") -> list[list[str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f, delimiter=delimiter))


def dir_listing(path: Path) -> list[str]:
    return sorted(p.name for p in path.iterdir())


_LEX = re.compile(rb'"(?:[^"\\]|\\.)*"|-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?|true|false|null'
                  rb'|[{}\[\],:]|[ \t\r\n]+')


def number_tokens(data: bytes) -> list[bytes]:
    """Number tokens outside strings, in document order (``data`` must be valid JSON text)."""
    out = []
    pos = 0
    for m in _LEX.finditer(data):
        assert m.start() == pos, data[pos:pos + 40]
        pos = m.end()
        tok = m.group()
        if tok[:1] in b"-0123456789":
            out.append(tok)
    assert pos == len(data)
    return out


# ----------------------------------------------------------------------------------
# help, version, usage errors (argparse localized)
# ----------------------------------------------------------------------------------
def test_help_is_uzbek_by_default(run) -> None:
    r = run("--help")
    assert r.code == 0
    assert r.out.startswith("foydalanish: jsonreshape")
    for text in ("parametrlar:", "buyruqlar:", "Misollar:", "Chiqish kodlari", "shu yordamni ko'rsatish",
                 "inspect", "search", "filter", "sort", "convert"):
        assert text in r.out, text
    assert "usage:" not in r.out and "show this help" not in r.out


def test_help_in_english_and_per_command(run) -> None:
    r = run("--lang", "en", "--help")
    assert r.code == 0 and r.out.startswith("usage: jsonreshape")
    assert "options:" in r.out and "commands:" in r.out and "Exit codes" in r.out
    for command, needle in (("inspect", "--sample"), ("search", "--translit"), ("filter", "--jmespath"),
                            ("sort", "--memory-limit"), ("convert", "--sort-by")):
        r = run(command, "--help", "--lang", "en")
        assert r.code == 0, command
        assert r.out.startswith(f"usage: jsonreshape {command}"), command
        assert needle in r.out and "input options:" in r.out and "general options:" in r.out
        r = run(command, "--help")
        assert r.out.startswith(f"foydalanish: jsonreshape {command}") and "kirish fayli parametrlari:" in r.out


def test_argparse_hooks_are_restored(run) -> None:
    before = (argparse._, argparse.ngettext)  # type: ignore[attr-defined]
    run("--help")
    run("filter")
    assert (argparse._, argparse.ngettext) == before  # type: ignore[attr-defined]


def test_version(run) -> None:
    from jsonreshape import core

    r = run("--version")
    assert r.code == 0
    lines = r.out.splitlines()
    assert lines[0] == f"jsonreshape {VERSION}"
    assert lines[1].startswith("Yadro: ")
    expected = {"rust": "Rust (tezkor)", "python": "sof Python (sekin)"}[core.BACKEND]
    assert expected in lines[1]
    assert lines[2].startswith("Python: ") and "64-bit" in lines[2]
    r = run("--lang", "en", "--version")
    assert "Core: " in r.out


@pytest.mark.parametrize("argv, needles", [
    ([], ["foydalanish: jsonreshape", "jsonreshape: xato: buyruq ko'rsatilmagan"]),
    (["frobnicate"], ["noto'g'ri tanlov: 'frobnicate'"]),
    (["inspect"], ["quyidagi argumentlar majburiy: INPUT"]),
    (["filter", "a.json", "b.json"], ["--where --jmespath --preset argumentlaridan bittasi majburiy"]),
    (["filter", "a.json", "b.json", "--where", "x", "--jmespath", "y"],
     ["--jmespath argumenti: --where argumenti bilan birga ishlatib bo'lmaydi"]),
    (["sort", "a.json", "b.json"], ["quyidagi argumentlar majburiy: --by"]),
    (["sort", "a.json", "b.json", "--by"], ["--by argumenti: bitta qiymat kutilgan"]),
    (["inspect", "a.json", "--format", "xml"], ["--format argumenti: noto'g'ri tanlov: 'xml'"]),
    (["inspect", "a.json", "--sample", "0"], ["musbat butun son bo'lishi kerak, berilgan: 0"]),
    (["search", "a.json", "x", "--limit", "-1"], ["0 yoki undan katta butun son bo'lishi kerak, berilgan: -1"]),
    (["convert", "a.json", "b.json", "--indent", "99"], ["0 dan 16 gacha son"]),
    (["sort", "a.json", "b.json", "--by", "id", "--memory-limit", "lots"], ["hajm noto'g'ri: lots"]),
    (["inspect", "a.json", "--bogus"], ["foydalanish: jsonreshape inspect", "tanilmagan argumentlar: --bogus"]),
    (["search", "a.json", "x", "--keys", "--values"], ["--keys argumenti bilan birga ishlatib bo'lmaydi"]),
])
def test_usage_errors_exit_1_with_localized_argparse_text(run, argv: list[str], needles: list[str]) -> None:
    r = run(*argv)
    assert r.code == 1, r.err
    assert r.out == ""
    for needle in needles:
        assert needle in r.err, (needle, r.err)
    assert "Yordam: jsonreshape" in r.err
    for english in ("usage:", "error:", "required", "invalid choice", "unrecognized", "expected one"):
        assert english not in r.err


def test_usage_error_in_english(run) -> None:
    r = run("--lang", "en", "sort", "a.json", "b.json")
    assert r.code == 1
    assert "usage: jsonreshape sort" in r.err and "the following arguments are required: --by" in r.err
    r = run("inspect", "a.json", "--format", "xml", "--lang", "en")
    assert "invalid choice: 'xml'" in r.err and "Help: jsonreshape inspect --help" in r.err


def test_argparse_message_table_covers_this_python() -> None:
    """Every user-facing gettext msgid of the running argparse is translated, and every
    translation key exists in both languages."""
    tree = ast.parse(Path(argparse.__file__).read_text(encoding="utf-8"))
    msgids = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in ("_", "ngettext"):
            for arg in node.args[:2] if node.func.id == "ngettext" else node.args[:1]:
                if isinstance(arg, ast.Constant) and isinstance(arg.value, str):
                    msgids.add(arg.value)
    developer_errors = {'argument "-" with mode %r', "conflicting option string: %s",
                        "conflicting option strings: %s"}
    missing = sorted(m for m in msgids - developer_errors if m not in cli.ARGPARSE_MESSAGES)
    assert not missing, missing
    for key in set(cli.ARGPARSE_MESSAGES.values()):
        entry = i18n.MESSAGES[key]
        assert entry["uz"] and entry["en"], key
        # %-placeholders must agree between the languages
        pattern = re.compile(r"%(?:\([a-z_]+\))?[sr]")
        assert sorted(pattern.findall(entry["uz"])) == sorted(pattern.findall(entry["en"])), key


def test_dynamic_cli_keys_exist() -> None:
    for name in ("objects", "mixed", "values", "empty"):
        assert i18n.has(f"cli.inspect.items.{name}")
    for mode in ("bytes", "parse"):
        assert i18n.has(f"search.mode.{mode}")
    from jsonreshape import search, writer
    from jsonreshape.reader import KINDS

    assert list(cli.INPUT_FORMATS) == ["auto", *KINDS]

    assert set(cli.OUTPUT_FORMATS) == set(writer.FORMATS)
    assert set(cli.HIT_EXPORT_FORMATS) == set(search.EXPORT_FORMATS)


def test_dash_values_after_by_are_joined() -> None:
    assert cli._join_dash_values(["sort", "a", "b", "--by", "-id", "--by", "name"]) == \
        ["sort", "a", "b", "--by=-id", "--by", "name"]
    assert cli._join_dash_values(["convert", "a", "b", "--sort-by", "-x.y"]) == ["convert", "a", "b", "--sort-by=-x.y"]
    # a real option is not swallowed, and nothing changes after "--"
    assert cli._join_dash_values(["sort", "a", "b", "--by", "--compact"]) == ["sort", "a", "b", "--by", "--compact"]
    assert cli._join_dash_values(["search", "--", "--by", "-x"]) == ["search", "--", "--by", "-x"]


# ----------------------------------------------------------------------------------
# inspect
# ----------------------------------------------------------------------------------
def test_inspect_report_uz(run, src: Path) -> None:
    from jsonreshape import core

    r = run("inspect", src)
    assert r.code == 0, r.err
    out = r.out
    assert re.search(r"Fayl:\s+" + re.escape(str(src)) + "\n", out)
    size = src.stat().st_size
    assert re.search(r"Hajm:\s+[\d.]+ [KM]?B \(" + re.escape(f"{size:,}".replace(",", " ")) + r" bayt\)\n", out)
    assert re.search(r"Kodirovka:\s+UTF-8\n", out)
    assert re.search(r"Format:\s+JSON massiv \(avtomatik\)\n", out)
    assert re.search(r"Yozuvlar soni:\s+8\n", out)
    assert re.search(r"Indeks:\s+\d+\.\d\d s, \d+(\.\d)? MB/s\n", out)
    assert {"rust": "Rust (tezkor)", "python": "sof Python (sekin)"}[core.BACKEND] in out
    assert re.search(r"Satrlarda:\s+\\ bilan yozilgan belgilar: yo'q; \\uXXXX kodlari: yo'q; "
                     r"ASCII bo'lmagan matn: ha\n", out)
    # the schema table and its summary follow the key/value block
    assert re.search(r"Yo'l\s+Turlar\s+Mavjudlik\s+Misollar", out)
    assert re.search(r"\ncontact\.address\.city\s+satr 100%", out)
    assert "Barcha yozuvlar tahlil qilindi: 8 ta" in out
    # second run: the index comes from the cache; --no-cache rebuilds it
    r = run("inspect", src)
    assert re.search(r"Indeks:\s+keshdan\n", r.out)
    r = run("inspect", src, "--no-cache")
    assert "keshdan" not in r.out


def test_inspect_report_en_and_sample(run, rnd: Path) -> None:
    r = run("--lang", "en", "inspect", rnd, "--sample", "5")
    assert r.code == 0, r.err
    assert re.search(r"Format:\s+JSON array \(detected\)\n", r.out)
    assert re.search(r"Records:\s+60\n", r.out)
    assert "Sample: 5 of 60 records analysed" in r.out


def test_inspect_json_has_exact_values(run, src: Path) -> None:
    from jsonreshape import core

    r = run("inspect", src, "--json", "--sample", "3")
    assert r.code == 0, r.err
    data = json.loads(r.out)
    assert data["file"] == str(src.resolve()) or Path(data["file"]) == src.resolve()
    assert data["size"] == src.stat().st_size == data["data_size"]
    assert data["encoding"] == "utf-8" and data["compression"] is None
    assert data["format"]["kind"] == "array" and data["format"]["auto"] is True
    assert data["format"]["label"] == "JSON massiv"
    assert data["records"] == 8
    assert data["index"]["complete"] is True and data["index"]["from_cache"] is False
    assert isinstance(data["index"]["seconds"], float) and data["index"]["bytes"] > 0
    assert data["core_backend"] == core.BACKEND
    assert data["flags"] == {"escapes": False, "u_escapes": False, "non_ascii": True}
    assert data["error"] is None and data["broken_lines"] == {"count": 0, "first": []}
    assert data["schema"]["scanned"] == 3 and data["schema"]["total_records"] == 8
    assert any(f["text"] == "orders[*].sum" for f in data["schema"]["fields"])


def test_inspect_wrapped_keyed_and_forced_formats(run, tmp_path: Path) -> None:
    recs = jsongen.records(12, seed=2)
    w = tmp_path / "w.json"
    jsongen.write_file(w, recs, "wrapped")
    r = run("inspect", w)
    assert r.code == 0, r.err
    assert re.search(r"Format:\s+Obyekt ichidagi massiv \(avtomatik\)\n", r.out)
    assert re.search(r"Yozuvlar yo'li:\s+data\n", r.out)
    assert re.search(r"Nomzod massivlar:\s+data: 12 ta element \(obyektlar\)", r.out)
    assert "Yozuvlar yo'li avtomatik tanlandi: data" in r.out
    data = json.loads(run("inspect", w, "--json").out)
    assert data["format"]["record_path"] == ["data"] and data["records"] == 12
    assert [c["text"] for c in data["format"]["candidates"]][0] == "data"
    # forced record path / format
    r = run("inspect", w, "--record-path", "meta.list", "--json")
    data = json.loads(r.out)
    assert r.code == 0 and data["records"] == 3 and data["format"]["auto"] is False
    k = tmp_path / "k.json"
    jsongen.write_file(k, recs, "keyed")
    data = json.loads(run("inspect", k, "--json").out)
    assert data["format"]["kind"] == "keyed" and data["format"]["member_records"] is True
    lines = tmp_path / "lines.txt"
    jsongen.write_file(lines, recs, "jsonl")
    data = json.loads(run("inspect", lines, "--format", "jsonl", "--json").out)
    assert data["format"]["kind"] == "jsonl" and data["records"] == 12
    r = run("inspect", lines, "--format", "jsonl")
    assert re.search(r"Format:\s+JSON Lines \(majburiy\)\n", r.out)


def test_inspect_gzip_and_utf16(run, tmp_path: Path) -> None:
    recs = jsongen.records(10, seed=4)
    gz = tmp_path / "data.json.gz"
    jsongen.write_file(gz, recs, "array", gz=True)
    r = run("inspect", gz)
    assert r.code == 0, r.err
    assert re.search(r"Kodirovka:\s+UTF-8, gzip\n", r.out) and "ochilgandan keyin" in r.out
    u16 = tmp_path / "u16.json"
    jsongen.write_file(u16, recs, "array", encoding="utf-16-le", bom=True)
    r = run("--lang", "en", "inspect", u16)
    assert re.search(r"Encoding:\s+UTF-16 LE\n", r.out) and re.search(r"Records:\s+10\n", r.out)


def test_inspect_jsonl_broken_lines(run, tmp_path: Path) -> None:
    p = tmp_path / "broken.jsonl"
    p.write_bytes(b'{"id": 1}\n{"id": 2,,}\n{"id": 3}\n{bad\n{"id": 5}\n')
    r = run("inspect", p)
    assert r.code == 0, r.err
    assert re.search(r"Yozuvlar soni:\s+3\n", r.out)
    assert re.search(r"Buzuq qatorlar:\s+2 ta\n", r.out)
    assert "2-qator, 10-ustun: kutilmagan belgi" in r.out and '{"id": 2,,}' in r.out
    assert "4-qator, 2-ustun: kutilmagan belgi" in r.out
    data = json.loads(run("inspect", p, "--json").out)
    assert data["broken_lines"]["count"] == 2
    assert [b["line"] for b in data["broken_lines"]["first"]] == [2, 4]


def test_inspect_json_error_exit_2_and_allow_partial(run, tmp_path: Path) -> None:
    p = tmp_path / "bad.json"
    p.write_bytes(b'[{"id": 1, "name": "Ali"},\n {"id": 2, "name": "Vali"}\n {"id": 3}]\n')
    r = run("inspect", p)
    assert r.code == 2
    assert re.search(r"Yozuvlar soni:\s+2 \(xatogacha\)\n", r.out)
    assert "kutilmagan belgi; ',' yoki ']' kutilgan (3-qator, 2-ustun, 55-bayt)" in r.out
    assert "Yo'l  " not in r.out  # no schema without --allow-partial
    assert "Xato: faylda JSON xatosi bor" in r.err and "--allow-partial" in r.err
    data = json.loads(run("inspect", p, "--json").out)
    assert data["error"]["line"] == 3 and data["error"]["column"] == 2 and data["error"]["records_before"] == 2
    assert data["schema"] is None
    r = run("inspect", p, "--allow-partial")
    assert r.code == 0 and "Barcha yozuvlar tahlil qilindi: 2 ta" in r.out


# ----------------------------------------------------------------------------------
# JSON errors, missing files, OUTPUT == INPUT
# ----------------------------------------------------------------------------------
@pytest.mark.parametrize("command", [["filter", "--where", "id > 0"], ["sort", "--by", "id"], ["convert"]])
def test_json_error_exit_2_with_caret_and_allow_partial(run, tmp_path: Path, command: list[str]) -> None:
    p = tmp_path / "bad.json"
    p.write_bytes(b'[{"id": 1, "name": "Ali"},\n {"id": 2, "name": "Vali"}\n {"id": 3}]\n')
    out = tmp_path / "out.json"
    r = run(command[0], p, out, *command[1:])
    assert r.code == 2
    lines = r.err.splitlines()
    assert lines[0] == ("Xato: JSON xatosi: kutilmagan belgi; ',' yoki ']' kutilgan "
                        "(3-qator, 2-ustun, 55-bayt).")
    context, caret = lines[1], lines[2]
    assert caret.strip() == "^" and context[caret.index("^")] == "{"
    assert context[:caret.index("^")].endswith('"Vali"}  ')
    assert "Maslahat: Xatogacha bo'lgan 2 ta yozuv" in r.err and "--allow-partial" in r.err
    assert "repair" not in r.err and "keyingi bosqich" not in r.err  # no development-plan promises
    assert not out.exists()
    r = run(command[0], p, out, *command[1:], "--allow-partial")
    assert r.code == 0, r.err
    assert "Ogohlantirish: Faylda JSON xatosi bor (3-qator, 2-ustun)" in r.err
    assert r.err.count("JSON xatosi") == 1  # the writer's own partial warning is not repeated
    assert ids(read_json(out)) == [1, 2]


def test_json_error_in_english(run, tmp_path: Path) -> None:
    p = tmp_path / "bad.json"
    p.write_bytes(b'[1, 2, }')
    r = run("--lang", "en", "convert", p, tmp_path / "o.jsonl")
    assert r.code == 2
    assert r.err.startswith("Error: JSON error: unexpected character; a value was expected (line 1, column 8, byte 7).")
    assert "Hint: Add --allow-partial to work with the 2 records before the error." in r.err


def test_missing_file_and_directory_exit_1(run, tmp_path: Path) -> None:
    r = run("inspect", tmp_path / "nope.json")
    assert r.code == 1 and r.err.startswith("Xato: Fayl topilmadi: ") and "nope.json" in r.err
    r = run("--lang", "en", "filter", tmp_path / "nope.json", tmp_path / "o.json", "--where", "x")
    assert r.code == 1 and r.err.startswith("Error: File not found: ")
    r = run("inspect", tmp_path)
    assert r.code == 1 and "Bu fayl emas, papka" in r.err


def test_empty_file_is_an_input_error(run, tmp_path: Path) -> None:
    p = tmp_path / "empty.json"
    p.write_bytes(b"  \n")
    r = run("inspect", p)
    assert r.code == 2 and "Fayl bo'sh" in r.err


@pytest.mark.parametrize("variant", ["same", "relative", "case", "dotted"])
def test_output_equal_to_input_is_refused(run, src: Path, monkeypatch: pytest.MonkeyPatch, variant: str) -> None:
    before = (src.read_bytes(), src.stat().st_mtime_ns)
    monkeypatch.chdir(src.parent)
    target = {"same": str(src), "relative": src.name, "case": str(src).upper() if sys.platform == "win32" else str(src),
              "dotted": str(src.parent) + os.sep + "." + os.sep + src.name}[variant]
    for argv in (["convert", src, target], ["filter", src, target, "--where", "id > 0"],
                 ["sort", src, target, "--by", "id"], ["convert", src, target, "--to", "csv"]):
        r = run(*argv)
        assert r.code == 1, (argv, r.err)
        assert "Natija manba faylning o'ziga yozilmaydi" in r.err and "Asl fayl hech qachon o'zgartirilmaydi" in r.err
    assert (src.read_bytes(), src.stat().st_mtime_ns) == before
    assert dir_listing(src.parent) == [src.name]


def test_search_export_to_input_is_refused(run, src: Path) -> None:
    before = src.read_bytes()
    r = run("search", src, "Anvar", "--export", src.with_suffix(".json"))
    assert r.code == 1 and "Natija manba faylning o'ziga yozilmaydi" in r.err
    assert src.read_bytes() == before


def test_output_folder_must_exist(run, src: Path, tmp_path: Path) -> None:
    r = run("convert", src, tmp_path / "no" / "such" / "dir" / "o.json")
    assert r.code == 1 and "Natija papkasi topilmadi" in r.err
    r = run("convert", src, tmp_path)
    assert r.code == 1 and "papka" in r.err


# ----------------------------------------------------------------------------------
# filter
# ----------------------------------------------------------------------------------
def test_filter_where_json(run, src: Path, tmp_path: Path) -> None:
    out = tmp_path / "o.json"
    r = run("filter", src, out, "--where", "score > 5")
    assert r.code == 0, r.err
    assert ids(read_json(out)) == [1, 2, 4, 7]
    assert r.out == ""
    assert re.fullmatch(r"4 / 8 yozuv -> .*o\.json \([\d.]+ [KM]?B, \d+\.\d s\)\n", r.err), r.err
    expected = [rec for rec in people() if rec["id"] in (1, 2, 4, 7)]
    assert canon(read_json(out)) == canon(expected)
    assert out.read_bytes().startswith(b"[\n  {\n    \"id\": 1,")  # indent 2 by default
    r = run("filter", src, out, "--where", 'contact.address.city == "Toshkent"', "-q")
    assert r.code == 0 and r.err == "" and ids(read_json(out)) == [1]


def test_filter_jmespath_and_bare_word_warning(run, src: Path, tmp_path: Path) -> None:
    out = tmp_path / "o.jsonl"
    r = run("filter", src, out, "--jmespath", "active")
    assert r.code == 0, r.err
    assert ids(read_jsonl(out)) == [1, 3, 4, 6, 7]
    r = run("filter", src, out, "--jmespath", "tags[1]")
    assert ids(read_jsonl(out)) == [1]
    r = run("filter", src, out, "--where", "name == Anvar")
    assert r.code == 0
    assert "Ogohlantirish: 'Anvar' maydon nomi deb tushunildi" in r.err


@pytest.mark.parametrize("argv, pos", [
    (["--where", "id < 5 and"], 10),
    (["--where", "score >> 3"], 7),
    (["--where", "contact.address.city == 'Toshkent"], 24),
    (["--jmespath", "foo[?"], 5),
])
def test_filter_bad_expression_exit_1_with_caret(run, src: Path, tmp_path: Path, argv: list[str], pos: int) -> None:
    out = tmp_path / "o.json"
    r = run("filter", src, out, *argv)
    assert r.code == 1
    lines = r.err.splitlines()
    assert lines[0].startswith("Xato: ")
    assert lines[1] == "  " + argv[1]
    assert lines[2] == "  " + " " * pos + "^"
    assert lines[3].startswith("Maslahat: ")
    assert not out.exists()


def test_filter_bad_expression_english(run, src: Path, tmp_path: Path) -> None:
    r = run("--lang", "en", "filter", src, tmp_path / "o.json", "--where", "id <")
    assert r.code == 1 and r.err.startswith("Error: Expression error at character 5:")


def test_filter_to_csv_xlsx_sqlite(run, src: Path, tmp_path: Path) -> None:
    import openpyxl

    csv_out = tmp_path / "o.csv"
    r = run("filter", src, csv_out, "--where", "active", "--columns", "id,name,contact.address.city,orders[*].oid")
    assert r.code == 0, r.err
    rows = read_csv(csv_out)
    assert rows[0] == ["id", "name", "contact.address.city", "orders[*].oid"]
    assert rows[1:] == [["1", "Anvar", "Toshkent", "7;8"], ["3", "dilmurod", "Samarqand", ""],
                        ["4", "Eldor", "O" + OKINA + "zbekiston", "9"], ["6", "child2", "", ""],
                        ["7", "Farrux", "toshkent", ""]]
    assert csv_out.read_bytes().startswith(b"\xef\xbb\xbf")
    xlsx_out = tmp_path / "o.xlsx"
    r = run("filter", src, xlsx_out, "--where", "id <= 2", "--columns", "id,name,score")
    assert r.code == 0, r.err
    wb = openpyxl.load_workbook(xlsx_out, read_only=True)
    values = [list(row) for row in wb.active.iter_rows(values_only=True)]
    wb.close()
    assert values == [["id", "name", "score"], [1, "Anvar", 10], [2, "Bobur", 25]]
    db = tmp_path / "o.sqlite"
    r = run("filter", src, db, "--where", "len(orders) >= 1")
    assert r.code == 0, r.err
    with sqlite3.connect(db) as con:
        assert [row[0] for row in con.execute('SELECT "id" FROM records ORDER BY _id')] == [1, 4]
        sums = [row[0] for row in con.execute('SELECT "sum" FROM orders ORDER BY _id')]
    con.close()
    # writer docstring "SQLite": numbers a double holds exactly are REAL (1.10 -> 1.1), integers
    # beyond 64 bits keep their exact text
    assert sums == [1.1, 0.1, "12345678901234567890"]


def test_filter_on_keyed_and_wrapped_documents(run, tmp_path: Path) -> None:
    recs = people()
    k = tmp_path / "k.json"
    jsongen.write_file(k, recs, "keyed", keys=[f"u{r['id']}" for r in recs])
    out = tmp_path / "k_out.json"
    r = run("filter", k, out, "--where", '$key in ["u2", "u5"] or id == 7')
    assert r.code == 0, r.err
    assert list(read_json(out)) == ["u2", "u5", "u7"]
    w = tmp_path / "w.json"
    jsongen.write_file(w, recs, "wrapped")
    out = tmp_path / "w_out.json"
    r = run("filter", w, out, "--where", "id > 6", "--compact")
    assert r.code == 0, r.err
    data = read_json(out)
    assert list(data) == ["meta", "data", "tail"] and ids(data["data"]) == [7, 8]
    assert out.read_bytes().startswith(b'{"meta":{')
    r = run("filter", w, out, "--where", "id > 6", "--no-wrapper")
    assert ids(read_json(out)) == [7, 8]


# ----------------------------------------------------------------------------------
# sort
# ----------------------------------------------------------------------------------
@pytest.mark.parametrize("options, expected", [
    (["--by", "name"], [1, 8, 2, 3, 4, 7, 6, 5]),
    (["--by", "name", "--collation", "simple"], [1, 8, 2, 6, 5, 3, 4, 7]),
    (["--by", "name", "--collation", "unicode"], [1, 8, 2, 6, 5, 3, 4, 7]),
    (["--by", "name", "--no-natural"], [1, 8, 2, 3, 4, 7, 5, 6]),
    (["--by", "name", "--case-sensitive"], [8, 1, 2, 3, 4, 7, 6, 5]),
    (["--by", "-name"], [5, 6, 7, 4, 3, 2, 8, 1]),
    (["--by", "score"], [5, 6, 1, 2, 7, 4, 3, 8]),
    (["--by", "score", "--nulls", "first"], [8, 3, 5, 6, 1, 2, 7, 4]),
    (["--by", "-score"], [4, 2, 7, 1, 6, 5, 3, 8]),
    (["--by=-score"], [4, 2, 7, 1, 6, 5, 3, 8]),
    (["--by", "active", "--by", "-id"], [8, 5, 2, 7, 6, 4, 3, 1]),
    # Latin before Cyrillic, "oʻ" after "z", lowercase before uppercase on a tie, missing last
    (["--by", "contact.address.city", "--by", "id"], [3, 7, 1, 4, 2, 5, 6, 8]),
    (["--by", "$index", "--by", "-id"], [1, 2, 3, 4, 5, 6, 7, 8]),
])
def test_sort_orders(run, src: Path, tmp_path: Path, options: list[str], expected: list[int]) -> None:
    out = tmp_path / "sorted.jsonl"
    r = run("sort", src, out, *options)
    assert r.code == 0, r.err
    assert ids(read_jsonl(out)) == expected
    assert re.fullmatch(r"8 / 8 yozuv -> .*sorted\.jsonl \([\d.]+ [KM]?B, \d+\.\d s\)\n", r.err), r.err


def test_sort_external_equals_in_memory(run, rnd: Path, tmp_path: Path) -> None:
    a, b = tmp_path / "mem.json", tmp_path / "ext.json"
    assert run("sort", rnd, a, "--by", "name", "--by", "-score").code == 0
    r = run("sort", rnd, b, "--by", "name", "--by", "-score", "--memory-limit", "1KB")
    assert r.code == 0, r.err
    assert a.read_bytes() == b.read_bytes()


def test_sort_bad_field_exit_1_with_caret(run, src: Path, tmp_path: Path) -> None:
    r = run("sort", src, tmp_path / "o.json", "--by", "orders[")
    assert r.code == 1
    lines = r.err.splitlines()
    assert lines[0].startswith("Xato: Yo'l noto'g'ri yozilgan (7-belgi)")
    assert lines[1:3] == ["  orders[", "        ^"]
    r = run("sort", src, tmp_path / "o.json", "--by", "**.x")
    assert r.code == 1 and "'**'" in r.err


# ----------------------------------------------------------------------------------
# convert and output options
# ----------------------------------------------------------------------------------
def test_convert_round_trip_json_jsonl_json(run, rnd: Path, tmp_path: Path) -> None:
    jl = tmp_path / "a.jsonl"
    back = tmp_path / "b.json"
    assert run("convert", rnd, jl).code == 0
    assert run("convert", jl, back).code == 0
    assert canon(read_json(back)) == canon(read_json(rnd))
    assert len(read_jsonl(jl)) == 60
    # --jsonl / --to override the extension
    other = tmp_path / "c.data"
    r = run("convert", rnd, other, "--jsonl")
    assert r.code == 0 and other.read_bytes() == jl.read_bytes()
    r = run("convert", rnd, other, "--to", "jsonl")
    assert r.code == 0 and other.read_bytes() == jl.read_bytes()


def test_convert_filter_and_sort_together(run, src: Path, tmp_path: Path) -> None:
    out = tmp_path / "o.csv"
    r = run("convert", src, out, "--where", "active == true", "--sort-by", "-score", "--columns", "id,score",
            "--nulls", "first")
    assert r.code == 0, r.err
    assert read_csv(out) == [["id", "score"], ["3", ""], ["4", "1e400"], ["7", "25"], ["1", "10"], ["6", "3"]]
    assert "5 / 8 yozuv" in r.err


@pytest.mark.parametrize("options, check", [
    (["--compact"], lambda b: b.startswith(b'[{"id":1,"name":"Anvar"') and b"\n" not in b.rstrip(b"\n")),
    (["--indent", "4"], lambda b: b.startswith(b'[\n    {\n        "id": 1,')),
    (["--indent", "tab"], lambda b: b.startswith(b'[\n\t{\n\t\t"id": 1,')),
    (["--sort-keys", "--compact"], lambda b: b.startswith(b'[{"active":true,"contact":{"address":{"city":"Toshkent"}}')),
])
def test_json_output_options(run, src: Path, tmp_path: Path, options: list[str], check) -> None:
    out = tmp_path / "o.json"
    r = run("convert", src, out, *options)
    assert r.code == 0, r.err
    data = out.read_bytes()
    assert check(data), data[:80]
    if "--sort-keys" in options:
        assert canon(read_json(out)) == codec.dumps(people(), sort_keys=True)
    else:
        assert canon(read_json(out)) == canon(people())


def test_csv_options(run, src: Path, tmp_path: Path) -> None:
    out = tmp_path / "o.csv"
    r = run("convert", src, out, "--columns", "id,tags", "--sep", " | ", "--delimiter", ";", "--no-bom",
            "--where", "id <= 2")
    assert r.code == 0, r.err
    assert not out.read_bytes().startswith(b"\xef\xbb\xbf")
    assert read_csv(out, ";") == [["id", "tags"], ["1", "x | y"], ["2", ""]]
    r = run("convert", src, out, "--columns", "id,orders[*].oid", "--array-mode", "rows", "--delimiter", "tab",
            "--where", "id <= 2")
    assert r.code == 0, r.err
    assert read_csv(out, "\t") == [["id", "orders[*].oid"], ["1", "7"], ["1", "8"], ["2", ""]]
    assert "3 ta qator" in r.err
    r = run("convert", src, out, "--columns", "id,orders[*].oid", "--explode", "orders", "--where", "id == 1")
    assert r.code == 0, r.err
    assert read_csv(out) == [["id", "orders[*].oid"], ["1", "7"], ["1", "8"]]


def test_columns_split_respects_brackets_and_quotes() -> None:
    assert cli._split_columns('a, b.c ,d["x,y"],e[0],`k,1`') == ["a", "b.c", 'd["x,y"]', "e[0]", "`k,1`"]
    assert cli._split_columns(" , ") == []


def test_xlsx_split_files_with_small_sheet_limit(run, src: Path, tmp_path: Path) -> None:
    import openpyxl
    from jsonreshape.appdirs import settings_file

    settings_file().write_text(json.dumps({"xlsx_max_rows": 3}), encoding="utf-8")
    out = tmp_path / "split.xlsx"
    r = run("convert", src, out, "--columns", "id", "--xlsx-split", "files")
    assert r.code == 0, r.err
    parts = sorted(tmp_path.glob("split_*.xlsx"))
    assert [p.name for p in parts] == ["split_1.xlsx", "split_2.xlsx", "split_3.xlsx", "split_4.xlsx"]
    assert "4 ta fayl" in r.err and "split_1.xlsx" in r.err
    got = []
    for p in parts:
        wb = openpyxl.load_workbook(p, read_only=True)
        got.extend(row[0] for row in wb.active.iter_rows(values_only=True, min_row=2))
        wb.close()
    assert got == [1, 2, 3, 4, 5, 6, 7, 8]


def test_inapplicable_options_are_ignored_with_a_warning(run, src: Path, tmp_path: Path) -> None:
    out = tmp_path / "o.json"
    r = run("convert", src, out, "--columns", "id", "--no-bom", "--compact")
    assert r.code == 0
    assert "Ogohlantirish: --columns, --no-bom: json formatida ishlatilmaydi" in r.err
    out = tmp_path / "o.csv"
    r = run("convert", src, out, "--indent", "4", "--sort-keys", "--columns", "id")
    assert "--indent, --sort-keys: csv formatida ishlatilmaydi" in r.err


def test_output_format_errors(run, src: Path, tmp_path: Path) -> None:
    r = run("convert", src, tmp_path / "o.unknown")
    assert r.code == 1 and "Natija formatini fayl nomidan aniqlab bo'lmadi" in r.err and "--to json" in r.err
    r = run("convert", src, tmp_path / "o.json", "--jsonl", "--to", "csv")
    assert r.code == 1 and "--jsonl va --to csv birga ishlatilmaydi" in r.err
    r = run("convert", src, tmp_path / "o.csv", "--delimiter", "ab")
    assert r.code == 1 and "Eksport sozlamasi noto'g'ri" in r.err
    r = run("convert", src, tmp_path / "o.csv", "--columns", "a[")
    assert r.code == 1 and "Ustun yo'li noto'g'ri" in r.err
    assert dir_listing(tmp_path) == ["people.json"]


# ----------------------------------------------------------------------------------
# search
# ----------------------------------------------------------------------------------
def test_search_prints_hits_with_markers_and_summary(run, src: Path) -> None:
    r = run("search", src, "Toshkent")
    assert r.code == 0, r.err
    assert r.out.splitlines() == [
        "#0  contact.address.city  [qiymat]  " + cli.MARK_OPEN + "Toshkent" + cli.MARK_CLOSE,
        "#4  tags[0]  [qiymat]  " + cli.MARK_OPEN + "Toshkent" + cli.MARK_CLOSE + " shahri",
        "#6  contact.address.city  [qiymat]  " + cli.MARK_OPEN + "toshkent" + cli.MARK_CLOSE,
    ]
    assert "\x1b[" not in r.out
    assert re.match(r"3 ta moslik, 3 ta yozuvda; \S+ \S+ ko'rib chiqildi: \d+\.\d s, \d+ MB/s, (bayt rejimi|"
                    r"har bir yozuvni o'qish rejimi)\.\n$", r.err), r.err


@pytest.mark.parametrize("argv, expected", [
    (["Toshkent", "--translit"], [(0, "contact.address.city"), (1, "contact.address.city"), (4, "tags[0]"),
                                  (6, "contact.address.city")]),
    (["Toshkent", "--case-sensitive"], [(0, "contact.address.city"), (4, "tags[0]")]),
    (["toshkent", "--whole-word", "--path", "contact.address.city"], [(0, "contact.address.city"),
                                                                       (6, "contact.address.city")]),
    (["O'zbekiston"], [(3, "contact.address.city")]),
    (["^child[0-9]+$", "--regex"], [(4, "name"), (5, "name")]),
    (["> 20", "--number"], [(1, "score"), (3, "score"), (3, "orders[0].sum"), (6, "score")]),
    (["score", "--keys", "--limit", "3"], [(0, "score"), (1, "score"), (2, "score")]),
    (["score", "--values"], []),
])
def test_search_modes(run, src: Path, argv: list[str], expected: list[tuple[int, str]]) -> None:
    r = run("search", src, *argv)
    assert r.code == 0, r.err
    got = [(int(line.split("  ")[0][1:]), line.split("  ")[1]) for line in r.out.splitlines()]
    assert got == expected


def test_search_keyed_document_key_hits(run, tmp_path: Path) -> None:
    recs = people()
    k = tmp_path / "k.json"
    jsongen.write_file(k, recs, "keyed", keys=[f"user{r['id']}" for r in recs])
    r = run("search", k, "user1", "--keys")
    assert r.out.splitlines() == ["#0  $  [kalit]  " + cli.MARK_OPEN + "user1" + cli.MARK_CLOSE]
    r = run("--lang", "en", "search", k, "user1", "--keys")
    assert r.out.splitlines() == ["#0  $  [key]  " + cli.MARK_OPEN + "user1" + cli.MARK_CLOSE]


def test_search_limit_count_and_max_hits(run, rnd: Path) -> None:
    from jsonreshape.reader import Document
    from jsonreshape.search import SearchOptions, reference_search

    with Document.open(rnd) as doc:
        total = len(reference_search(doc, SearchOptions.build("a")))
    assert total > 60
    r = run("search", rnd, "a")
    assert len(r.out.splitlines()) == 50
    assert f"Faqat dastlabki 50 ta moslik ko'rsatildi (jami {cli._num(total)} ta)" in r.err
    r = run("search", rnd, "a", "--limit", "5")
    assert len(r.out.splitlines()) == 5
    r = run("search", rnd, "a", "--limit", "0")
    assert r.out == "" and "Faqat dastlabki" not in r.err
    r = run("search", rnd, "a", "--max-hits", "7")
    assert len(r.out.splitlines()) == 7
    r = run("search", rnd, "a", "--count")
    assert r.code == 0
    assert re.fullmatch(rf"{total} ta moslik, \d+ ta yozuvda\n", r.out), r.out
    r = run("search", rnd, "a", "--count", "-q")
    assert re.fullmatch(rf"{total} ta moslik, \d+ ta yozuvda\n", r.out) and r.err == ""


def test_search_export_round_trip_csv_and_json(run, rnd: Path, tmp_path: Path) -> None:
    from jsonreshape.reader import Document
    from jsonreshape.search import SearchOptions, reference_search

    with Document.open(rnd) as doc:
        ref = reference_search(doc, SearchOptions.build("o"))
    assert len(ref) > 60
    csv_path = tmp_path / "hits.csv"
    r = run("search", rnd, "o", "--export", csv_path, "--limit", "3")
    assert r.code == 0, r.err
    assert len(r.out.splitlines()) == 3
    assert f"{cli._num(len(ref))} ta moslik yozildi" in r.err
    rows = read_csv(csv_path)
    assert rows[0] == ["record", "key", "path", "where", "text"]
    assert [tuple(row) for row in rows[1:]] == [(str(h.record), "", h.path_text, h.where, h.text) for h in ref]
    json_path = tmp_path / "hits.json"
    r = run("search", rnd, "o", "--export", json_path, "--count")
    assert r.code == 0 and r.out.startswith(f"{len(ref)} ta moslik")
    data = json.loads(json_path.read_text(encoding="utf-8"))
    assert [(d["record"], d["path"], d["where"], d["text"], [tuple(s) for s in d["spans"]]) for d in data] == \
        [(h.record, h.path_text, h.where, h.text, list(h.spans)) for h in ref]
    r = run("search", rnd, "o", "--export", tmp_path / "hits.txt")
    assert r.code == 1 and "Qidiruv natijalarini bu formatda saqlab bo'lmaydi: txt" in r.err


def test_search_bad_regex_number_and_path_exit_1_with_caret(run, src: Path) -> None:
    r = run("search", src, "ab(c", "--regex")
    assert r.code == 1
    lines = r.err.splitlines()
    assert lines[0].startswith("Xato: Regex ifodasida xato (3-belgi)")
    assert lines[1:3] == ["  ab(c", "    ^"]
    r = run("search", src, "> x", "--number")
    assert r.code == 1 and r.err.startswith("Xato: Son shartida xato")
    r = run("search", src, "a", "--path", "a[")
    assert r.code == 1 and "Yo'l noto'g'ri yozilgan (2-belgi)" in r.err and "\n  a[\n   ^\n" in r.err
    r = run("search", src, "")
    assert r.code == 1 and "Qidiruv matni bo'sh" in r.err


def test_search_streams_hits_through_on_hits(run, src: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape import search as search_mod

    original = search_mod.search
    seen = {}

    def spy(*args, **kwargs):
        seen["on_hits"] = kwargs.get("on_hits")
        seen["max_hits"] = kwargs.get("max_hits")
        return original(*args, **kwargs)

    monkeypatch.setattr(search_mod, "search", spy)
    r = run("search", src, "Toshkent", "--limit", "2")
    assert r.code == 0 and len(r.out.splitlines()) == 2
    assert callable(seen["on_hits"]) and seen["max_hits"] == 2


def test_search_history_is_recorded(run, src: Path) -> None:
    from jsonreshape.search import SearchHistory

    run("search", src, "Toshkent", "--translit", "--keys")
    run("search", src, "Bobur")
    entries = SearchHistory().list()
    assert [e["text"] for e in entries[:2]] == ["Bobur", "Toshkent"]
    assert entries[1]["options"]["translit"] is True and entries[1]["options"]["scope"] == "keys"


def test_search_highlighting_on_a_terminal(run, src: Path, monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    class TTY(io.StringIO):
        def isatty(self) -> bool:
            return True

    fake = TTY()
    monkeypatch.setattr(sys, "stdout", fake)
    monkeypatch.setattr(cli, "_enable_ansi", lambda: True)
    assert cli.main(["search", str(src), "Toshkent", "--limit", "1"]) == 0
    assert fake.getvalue() == ("#0  contact.address.city  [qiymat]  " + cli.HIGHLIGHT_ON + "Toshkent"
                               + cli.HIGHLIGHT_OFF + "\n")
    fake = TTY()
    monkeypatch.setattr(sys, "stdout", fake)
    monkeypatch.setenv("NO_COLOR", "1")
    assert cli.main(["search", str(src), "Toshkent", "--limit", "1"]) == 0
    assert "\x1b[" not in fake.getvalue() and cli.MARK_OPEN + "Toshkent" + cli.MARK_CLOSE in fake.getvalue()
    capsys.readouterr()


def test_display_escapes_control_characters() -> None:
    text = "a\nb\tc" + chr(7) + "d" + chr(0x2028) + "e" + chr(0xD800)
    shown = cli._display_text(text)
    esc = "\\" + "u"
    assert shown == "a\\nb\\tc" + esc + "0007d" + esc + "2028e" + esc + "d800"
    assert cli._highlight("xx\nyy", [(1, 4)], False) == "x" + cli.MARK_OPEN + "x\\ny" + cli.MARK_CLOSE + "y"
    assert cli._highlight("abc", [(5, 9), (0, 1)], True) == cli.HIGHLIGHT_ON + "a" + cli.HIGHLIGHT_OFF + "bc"


# ----------------------------------------------------------------------------------
# cancellation (exit code 3, nothing partial left behind)
# ----------------------------------------------------------------------------------
class _CountdownEvent(threading.Event):
    """An Event that reports "set" from the n-th is_set() call on -- counting only the
    calls made once ``ready()`` is true (default: all of them)."""

    def __init__(self, n: int, ready: Callable[[], bool] | None = None) -> None:
        super().__init__()
        self.calls = 0
        self.n = n
        self.ready = ready

    def is_set(self) -> bool:  # noqa: D401
        if self.ready is None or self.ready():
            self.calls += 1
        return self.calls >= self.n or super().is_set()


def _big_people(tmp_path: Path, count: int = 5000) -> Path:
    p = tmp_path / "big.json"
    rec = people()
    jsongen.write_file(p, [dict(rec[i % len(rec)], id=i) for i in range(count)], "array")
    return p


def test_cancelled_before_start_exit_3(run, src: Path, tmp_path: Path) -> None:
    token = CancelToken()
    token.cancel()
    out = tmp_path / "o.json"
    r = run("convert", src, out, cancel=token)
    assert r.code == 3
    assert r.err.strip() == "Amal bekor qilindi."
    assert not out.exists()
    r = run("--lang", "en", "inspect", src, cancel=token)
    assert r.code == 3 and r.err.strip() == "The operation was cancelled."


@pytest.mark.parametrize("fmt, nth_check", [("json", 2), ("jsonl", 2), ("csv", 3), ("xlsx", 3), ("sqlite", 3)])
def test_cancel_during_export_leaves_no_partial_output(run, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                                       fmt: str, nth_check: int) -> None:
    """Cancelled at the n-th cancel check made once the writer has created its temporary
    file (however many checks it makes before that), i.e. in the middle of writing: exit 3
    and nothing left behind."""
    from jsonreshape import fileio, writer

    big = _big_people(tmp_path)
    token = CancelToken()
    original = writer.export
    original_temp = fileio._create_temp
    state: dict[str, Any] = {"temps": []}

    def export_then_cancel(*args, **kwargs):
        token._event = _CountdownEvent(nth_check, ready=lambda: bool(state["temps"]))
        state["event"] = token._event
        return original(*args, **kwargs)

    def spy_temp(*args, **kwargs):
        fd, path = original_temp(*args, **kwargs)
        state["temps"].append(path)
        return fd, path

    monkeypatch.setattr(writer, "export", export_then_cancel)
    monkeypatch.setattr(fileio, "_create_temp", spy_temp)
    out = tmp_path / f"o.{fmt}"
    columns = ["--columns", "id,name"] if fmt in ("csv", "xlsx") else []
    r = run("convert", big, out, *columns, cancel=token)
    assert r.code == 3, r.err
    assert r.err.strip().endswith("Amal bekor qilindi.")
    assert state["event"].calls >= nth_check
    assert state["temps"] and all(p.parent == tmp_path for p in state["temps"])  # it had started writing
    assert dir_listing(tmp_path) == ["big.json"]


def test_cancel_during_search_export(run, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape import search as search_mod

    big = _big_people(tmp_path)
    token = CancelToken()
    original = search_mod.export_hits

    def export_then_cancel(*args, **kwargs):
        token._event = _CountdownEvent(2)
        return original(*args, **kwargs)

    monkeypatch.setattr(search_mod, "export_hits", export_then_cancel)
    r = run("search", big, "Toshkent", "--export", tmp_path / "hits.csv", cancel=token)
    assert r.code == 3
    assert dir_listing(tmp_path) == ["big.json"]


def test_sigint_cancels_and_second_sigint_interrupts(run, src: Path, tmp_path: Path,
                                                     monkeypatch: pytest.MonkeyPatch) -> None:
    if threading.current_thread() is not threading.main_thread():  # pragma: no cover
        pytest.skip("signal handlers need the main thread")
    from jsonreshape import writer

    original = writer.export
    presses = {"n": 1}

    def export_with_ctrl_c(*args, **kwargs):
        for _ in range(presses["n"]):
            signal.raise_signal(signal.SIGINT)
        return original(*args, **kwargs)

    monkeypatch.setattr(writer, "export", export_with_ctrl_c)
    out = tmp_path / "o.json"
    r = run("convert", src, out)
    assert r.code == 3 and "Amal bekor qilindi." in r.err
    assert not out.exists()
    presses["n"] = 2  # the second Ctrl+C raises KeyboardInterrupt at once
    r = run("convert", src, out)
    assert r.code == 3 and "Amal bekor qilindi." in r.err
    assert not out.exists()


def test_keyboard_interrupt_is_exit_3(run, src: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape import recordset

    def boom(*args, **kwargs):
        raise KeyboardInterrupt

    monkeypatch.setattr(recordset, "sort_records", boom)
    r = run("sort", src, tmp_path / "o.json", "--by", "id")
    assert r.code == 3 and "Amal bekor qilindi." in r.err


# ----------------------------------------------------------------------------------
# internal and OS errors, logging
# ----------------------------------------------------------------------------------
def test_internal_error_exit_4_with_log_path(run, src: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape import log as applog
    from jsonreshape import schema

    def boom(*args, **kwargs):
        raise RuntimeError("kutilmagan holat")

    monkeypatch.setattr(schema, "infer_schema", boom)
    r = run("inspect", src)
    assert r.code == 4
    assert "Ichki xato: RuntimeError: kutilmagan holat" in r.err
    log_file = applog.log_path()
    assert f"Tafsilotlar log faylida: {log_file}" in r.err
    assert "Traceback" not in r.err
    text = log_file.read_text(encoding="utf-8")
    assert "internal error" in text and "Traceback" in text and "kutilmagan holat" in text
    r = run("-vv", "inspect", src)
    assert r.code == 4 and "Traceback (most recent call last)" in r.err
    r = run("--lang", "en", "inspect", src)
    assert "Internal error: RuntimeError" in r.err and "This is a bug in the program." in r.err


def test_os_error_is_a_user_error(run, src: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape import writer

    def denied(*args, **kwargs):
        raise PermissionError(13, "Access is denied", str(tmp_path / "o.json"))

    monkeypatch.setattr(writer, "export", denied)
    r = run("convert", src, tmp_path / "o.json")
    # a localized message with the path and a hint; the OS text only as a detail line
    assert r.code == 1 and r.err.startswith(f"Xato: Ruxsat yo'q: {tmp_path / 'o.json'}\n"), r.err
    assert "  tizim xabari: Access is denied\n" in r.err and "Maslahat: " in r.err

    def strange(*args, **kwargs):
        raise OSError(errno.EIO, "I/O error", str(tmp_path / "o.json"))

    monkeypatch.setattr(writer, "export", strange)
    r = run("convert", src, tmp_path / "o.json")
    assert r.code == 1 and r.err.startswith("Xato: fayl tizimi xatosi: I/O error: ")


def test_log_file_records_the_run(run, src: Path) -> None:
    from jsonreshape import log as applog

    assert run("inspect", src, "--sample", "2").code == 0
    text = applog.log_path().read_text(encoding="utf-8")
    assert f"jsonreshape {VERSION}: inspect {src} --sample 2" in text
    assert "exit code 0 after" in text and "environment: jsonreshape" in text


def test_verbose_shows_log_records_and_environment_warnings(run, src: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import ijson

    r = run("inspect", src)
    assert "INFO" not in r.err
    monkeypatch.setattr(ijson, "backend", "python")
    r = run("-v", "inspect", src)
    assert r.code == 0
    assert "INFO jsonreshape.cli: jsonreshape" in r.err
    assert "WARNING jsonreshape.env: ijson backend is 'python', expected 'yajl2_c'" in r.err
    r = run("inspect", src, "-vv")
    assert "DEBUG" in r.err or "INFO" in r.err
    # the console handler is removed again after the run
    assert not [h for h in logging.getLogger().handlers
                if isinstance(h, logging.StreamHandler) and not isinstance(h, logging.FileHandler)
                and getattr(h, "_jsonreshape_handler", False)]


# ----------------------------------------------------------------------------------
# language selection
# ----------------------------------------------------------------------------------
def test_language_priority(run, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape.appdirs import settings_file
    from jsonreshape.settings import _reset_for_tests

    missing = tmp_path / "missing.json"
    monkeypatch.delenv("JSONRESHAPE_LANG")
    assert run("inspect", missing).err.startswith("Xato: Fayl topilmadi")      # default: uz
    settings_file().write_text(json.dumps({"language": "en"}), encoding="utf-8")
    _reset_for_tests()
    assert run("inspect", missing).err.startswith("Error: File not found")     # settings
    monkeypatch.setenv("JSONRESHAPE_LANG", "uz")
    assert run("inspect", missing).err.startswith("Xato: Fayl topilmadi")      # env > settings
    assert run("inspect", missing, "--lang", "en").err.startswith("Error:")    # --lang > env
    assert run("--lang=en", "inspect", missing).err.startswith("Error:")
    assert run("--lang", "uz", "inspect", missing, "--lang", "en").err.startswith("Error:")


# ----------------------------------------------------------------------------------
# progress line
# ----------------------------------------------------------------------------------
class _TTY(io.StringIO):
    def isatty(self) -> bool:
        return True


def test_progress_line_renders_and_clears(monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _TTY()
    monkeypatch.setattr(sys, "stderr", fake)
    i18n.set_language("uz")
    line = cli._ProgressLine(enabled=True)
    prog = line.stage("Indekslanmoqda")
    assert prog is not None
    prog.set_total(2_000_000)
    prog.update(1_000_000)
    text = fake.getvalue()
    assert "\rIndekslanmoqda: " in text
    last = text.split("\r")[-1]
    assert "/ 2.0 MB" in last and "/s" in last and "qoldi" in last
    prog.finish()
    final = fake.getvalue()
    assert final.endswith("\r") and final.split("\r")[-2].strip() == ""
    assert cli._ProgressLine(enabled=False).stage("x") is None
    # unknown total
    line2 = cli._ProgressLine(enabled=True)
    p2 = line2.stage("Qidirilmoqda")
    p2.update(500)
    assert "Qidirilmoqda: 500 B" in fake.getvalue().split("\r")[-1]
    line2.clear()


def test_progress_only_on_a_terminal_and_not_with_quiet(src: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                                         capsys) -> None:
    captured_stderr = sys.stderr
    fake = _TTY()
    monkeypatch.setattr(sys, "stderr", fake)
    out = tmp_path / "o.json"
    assert cli.main(["convert", str(src), str(out), "--no-cache"]) == 0
    text = fake.getvalue()
    assert "\rIndekslanmoqda: " in text and "\rYozilmoqda: " in text
    # every progress line is cleared before the summary line is written
    summary = text.rsplit("\r", 1)[-1]
    assert summary.startswith("8 / 8 yozuv -> ")
    fake = _TTY()
    monkeypatch.setattr(sys, "stderr", fake)
    assert cli.main(["convert", str(src), str(out), "--no-cache", "-q"]) == 0
    assert fake.getvalue() == ""
    monkeypatch.setattr(sys, "stderr", captured_stderr)
    capsys.readouterr()
    # not a terminal (captured): no progress at all
    r_code = cli.main(["convert", str(src), str(out), "--no-cache"])
    captured = capsys.readouterr()
    assert r_code == 0 and "\r" not in captured.err


# ----------------------------------------------------------------------------------
# backends
# ----------------------------------------------------------------------------------
def test_backend_flag_in_a_running_process_warns(run, src: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jsonreshape import core

    other = "python" if core.BACKEND == "rust" else "rust"
    # the environment check would load the other backend's module; a python-backend test run
    # must never load the Rust extension
    monkeypatch.setattr(cli, "_check_environment", lambda: None)
    r = run("--backend", other, "inspect", src, "--sample", "1")
    assert r.code == 0
    assert f"--backend {other} faqat yangi ishga tushirishda ta'sir qiladi" in r.err
    assert os.environ["JSONRESHAPE_BACKEND"] == other
    r = run("--backend", core.BACKEND, "inspect", src, "--sample", "1")
    assert "--backend" not in r.err


_DRIVER = r'''
import contextlib, io, json, os, sys
from jsonreshape import cli
backend, workdir = sys.argv[1], sys.argv[2]
src = os.path.join(workdir, "data.json")
out = os.path.join(workdir, backend)
os.makedirs(out, exist_ok=True)
results = {}
runs = [
    ("inspect", ["inspect", src, "--json", "--no-cache"]),
    ("filter", ["filter", src, os.path.join(out, "filter.json"), "--where", "active == true and score > 0"]),
    ("jmespath", ["filter", src, os.path.join(out, "jmes.jsonl"), "--jmespath", "contact.address.city"]),
    ("sort", ["sort", src, os.path.join(out, "sort.jsonl"), "--by", "name", "--by", "-score"]),
    ("convert", ["convert", src, os.path.join(out, "conv.csv"), "--where", "id >= 10", "--sort-by", "-id"]),
    ("tab", ["convert", src, os.path.join(out, "conv.json"), "--indent", "tab", "--sort-keys"]),
    ("search", ["search", src, "Toshkent", "--limit", "1000"]),
    ("translit", ["search", src, "toshkent", "--translit", "--limit", "1000"]),
    ("regex", ["search", src, "^[0-9]+$", "--regex", "--limit", "1000"]),
    ("number", ["search", src, "> 100", "--number", "--limit", "1000"]),
]
for name, argv in runs:
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = cli.main(["--backend", backend, "-q"] + argv)
    results[name] = {"code": code, "stdout": buf.getvalue()}
from jsonreshape import core
results["core_backend"] = core.BACKEND
with open(os.path.join(out, "results.json"), "w", encoding="utf-8") as f:
    json.dump(results, f, ensure_ascii=False)
'''


def rust_extension_built() -> bool:
    """Is the Rust core present (without importing it into this process)?"""
    import importlib.util

    return importlib.util.find_spec("jsonreshape.core._rscore") is not None


@pytest.mark.slow
def test_backends_give_identical_outputs(tmp_path: Path) -> None:
    if not rust_extension_built():
        pytest.skip("the Rust core is not built")
    jsongen.write_file(tmp_path / "data.json", jsongen.records(150, seed=21), "array", style="pretty2")
    driver = tmp_path / "driver.py"
    driver.write_bytes(_DRIVER.encode("utf-8"))
    env = {k: v for k, v in os.environ.items() if k not in ("JSONRESHAPE_BACKEND", "PYTHONIOENCODING")}
    env["JSONRESHAPE_LANG"] = "uz"
    results = {}
    for backend in ("python", "rust"):
        proc = subprocess.run([sys.executable, str(driver), backend, str(tmp_path)], cwd=str(ROOT), env=env,
                              capture_output=True, timeout=300)
        assert proc.returncode == 0, proc.stderr.decode("utf-8", "replace")
        results[backend] = json.loads((tmp_path / backend / "results.json").read_text(encoding="utf-8"))
        assert results[backend]["core_backend"] == backend  # --backend really selected the core
    py, rs = results["python"], results["rust"]
    for name in ("filter", "jmespath", "sort", "convert", "tab", "search", "translit", "regex", "number"):
        assert py[name]["code"] == rs[name]["code"] == 0, name
        assert py[name]["stdout"] == rs[name]["stdout"], name
    assert py["search"]["stdout"] and py["number"]["stdout"]
    for fname in ("filter.json", "jmes.jsonl", "sort.jsonl", "conv.csv", "conv.json"):
        a = (tmp_path / "python" / fname).read_bytes()
        assert a == (tmp_path / "rust" / fname).read_bytes(), fname
        assert len(a) > 20, fname
    ip, ir = json.loads(py["inspect"]["stdout"]), json.loads(rs["inspect"]["stdout"])
    for d in (ip, ir):
        d.pop("core_backend")
        d["index"] = {k: v for k, v in d["index"].items() if k in ("complete", "from_cache", "bytes")}
        d["schema"].pop("elapsed")
        d.pop("schema_summary")
    assert ip == ir


# ----------------------------------------------------------------------------------
# AGENT_PROMPT §14 test 24 through the CLI
# ----------------------------------------------------------------------------------
PROMPT_TOKENS = ["0.1", "1e400", "12345678901234567890", "-0.0", "1.10"]
MORE_TOKENS = ["-0", "1E5", "2.5E+3", "123456789.123456789", "5e-324", "18446744073709551616"]
TOKENS = PROMPT_TOKENS + MORE_TOKENS


UZ_TEXT = "O" + OKINA + "zbekiston"


def _numbers_source(path: Path) -> None:
    recs = []
    for k, tok in enumerate(TOKENS):
        recs.append('{\n  "i": %d,\n  "n": %s,\n  "arr": [%s, %s],\n  "obj": {"x": %s, "t": "%s %d"}\n}'
                    % (k, tok, tok, tok, tok, UZ_TEXT, k))
    path.write_bytes(("[\n" + ",\n".join(recs) + "\n]\n").encode("utf-8"))


def _expected_tokens(order: list[int], *, sorted_keys: bool = False) -> list[bytes]:
    out = []
    for k in order:
        tok = TOKENS[k].encode()
        if sorted_keys:  # arr, i, n, obj
            out.extend([tok, tok, str(k).encode(), tok, tok])
        else:
            out.extend([str(k).encode(), tok, tok, tok, tok])
    return out


@pytest.mark.parametrize("argv, order, sorted_keys", [
    (["filter", "{out}.json", "--where", "i >= 0"], "all", False),
    (["filter", "{out}.jsonl", "--where", "i >= 0"], "all", False),
    (["filter", "{out}.json", "--where", "i >= 0", "--compact"], "all", False),
    (["sort", "{out}.json", "--by", "-i"], "reversed", False),
    (["sort", "{out}.jsonl", "--by", "-i"], "reversed", False),
    (["sort", "{out}.json", "--by", "n", "--indent", "tab"], "by_value", False),
    (["convert", "{out}.json"], "all", False),
    (["convert", "{out}.json", "--indent", "4"], "all", False),
    (["convert", "{out}.json", "--sort-keys"], "all", True),
    (["convert", "{out}.jsonl"], "all", False),
    (["convert", "{out}.jsonl", "--sort-keys"], "all", True),
    (["convert", "{out}.jsonl", "--where", "i >= 5", "--sort-by", "-i"], "five_down", False),
])
def test_number_tokens_survive_byte_identical(run, tmp_path: Path, argv: list[str], order: str,
                                             sorted_keys: bool) -> None:
    src = tmp_path / "numbers.json"
    _numbers_source(src)
    out = Path(argv[1].format(out=str(tmp_path / "out")))
    r = run(argv[0], src, out, *argv[2:])
    assert r.code == 0, r.err
    data = out.read_bytes()
    n = len(TOKENS)
    order_list = {
        "all": list(range(n)),
        "reversed": list(reversed(range(n))),
        "five_down": list(reversed(range(5, n))),
        "by_value": sorted(range(n), key=lambda k: (Decimal(TOKENS[k]), k)),
    }[order]
    assert number_tokens(data) == _expected_tokens(order_list, sorted_keys=sorted_keys)
    for tok in TOKENS:
        assert f'"{tok}"'.encode() not in data  # never turned into a string
    assert UZ_TEXT.encode("utf-8") in data and BACKSLASH_U not in data
